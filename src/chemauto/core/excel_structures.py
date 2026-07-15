from __future__ import annotations

import math
import tempfile
from pathlib import Path
from typing import Iterable, Sequence

import numpy as np
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from chemauto.core.gaussian_parser import LogSummary, delta_e_ev, minimum_etot

ATOM_COLORS = {
    "H": (210, 210, 210), "C": (75, 75, 75), "N": (45, 70, 200), "O": (190, 45, 45),
    "F": (70, 180, 70), "P": (235, 135, 35), "S": (220, 190, 45), "Cl": (50, 180, 60),
    "Fe": (165, 80, 45), "Co": (170, 75, 55), "Ni": (90, 150, 95), "Cu": (180, 105, 55),
    "Zn": (115, 120, 170), "As": (145, 80, 155), "Sb": (125, 70, 145), "I": (100, 55, 140),
}
COVALENT = {"H":0.31,"C":0.76,"N":0.71,"O":0.66,"F":0.57,"P":1.07,"S":1.05,"Cl":1.02,"Fe":1.32,"Co":1.26,"Ni":1.24,"Cu":1.32,"Zn":1.22,"As":1.19,"Sb":1.39,"I":1.39}
RADII = {"H":8,"C":13,"N":13,"O":14,"F":13,"Fe":22,"Co":22,"Ni":21,"Cu":21,"Zn":21,"As":18,"Sb":21,"I":20}


def pymol_available() -> bool:
    """Return True when the bundled/installed PyMOL backend can be imported.

    ChemAuto 5.3 used PyMOL for the Excel structure images.  ChemAuto 7.1.2
    bundles that backend again and keeps the in-house renderer as a fallback.
    """
    try:
        from pymol import cmd  # noqa: F401
        return True
    except Exception:
        return False


def parse_xyz(xyz: str) -> tuple[list[str], np.ndarray]:
    lines=[line for line in xyz.splitlines() if line.strip()]
    if len(lines)<3: raise ValueError("Invalid XYZ text")
    n=int(lines[0].strip()); body=lines[2:2+n]
    syms=[]; coords=[]
    for line in body:
        p=line.split()
        if len(p)<4: continue
        syms.append(p[0]); coords.append([float(p[1]),float(p[2]),float(p[3])])
    if len(syms)!=n: raise ValueError("XYZ atom count mismatch")
    return syms,np.asarray(coords,dtype=float)


def _project(coords: np.ndarray) -> tuple[np.ndarray,np.ndarray]:
    centered=coords-coords.mean(axis=0)
    if len(coords)>1:
        try:
            _,_,vh=np.linalg.svd(centered,full_matrices=False)
            projected=centered@vh[:2].T
            depth=centered@vh[2].T if vh.shape[0]>=3 else np.zeros(len(coords))
            return projected,depth
        except Exception: pass
    return centered[:,:2], centered[:,2] if centered.shape[1]>=3 else np.zeros(len(coords))


def render_xyz_png_qt(xyz: str, output: str | Path, size: int = 700) -> Path:
    """Built-in renderer implemented with Pillow (no PyMOL required).

    The historical function name is retained to avoid breaking callers from
    early Phase 3 builds. Rendering itself is GUI-independent and therefore
    also works in batch jobs and PyInstaller executables.
    """
    from PIL import Image, ImageDraw

    syms, coords = parse_xyz(xyz)
    pts, depth = _project(coords)
    minxy = pts.min(axis=0); maxxy = pts.max(axis=0)
    span = np.maximum(maxxy - minxy, 1e-6)
    margin = 70
    scale = min((size - 2 * margin) / span[0], (size - 2 * margin) / span[1])
    xy = (pts - (minxy + maxxy) / 2) * scale + size / 2
    xy[:, 1] = size - xy[:, 1]

    image = Image.new("RGB", (size, size), "white")
    draw = ImageDraw.Draw(image)
    bonds = []
    for i in range(len(syms)):
        for j in range(i + 1, len(syms)):
            cutoff = 1.25 * (COVALENT.get(syms[i], 0.8) + COVALENT.get(syms[j], 0.8))
            distance = float(np.linalg.norm(coords[i] - coords[j]))
            if 0.15 < distance <= cutoff:
                bonds.append((min(depth[i], depth[j]), i, j))
    for _, i, j in sorted(bonds):
        draw.line((float(xy[i,0]), float(xy[i,1]), float(xy[j,0]), float(xy[j,1])), fill=(55,55,55), width=max(6, size // 100))
    for _, i in sorted((depth[i], i) for i in range(len(syms))):
        radius = max(12, int(round(RADII.get(syms[i], 15) * 1.65)))
        x, y = float(xy[i,0]), float(xy[i,1])
        colour = ATOM_COLORS.get(syms[i], (110,110,110))
        draw.ellipse((x-radius, y-radius, x+radius, y+radius), fill=colour, outline=(30,30,30), width=2)
    out = Path(output); out.parent.mkdir(parents=True, exist_ok=True)
    image.save(out, format="PNG")
    return out


def render_xyz_png_pymol(xyz: str, output: str | Path) -> Path:
    try:
        from pymol import cmd
    except Exception as exc:
        raise RuntimeError(
            "The PyMOL renderer is unavailable. Use the ChemAuto built-in renderer "
            "or rebuild ChemAuto with the bundled PyMOL backend."
        ) from exc
    out=Path(output); xyz_path=out.with_suffix(".xyz"); xyz_path.write_text(xyz,encoding="utf-8")
    cmd.reinitialize(); cmd.load(str(xyz_path),"molecule"); cmd.hide("everything","molecule")
    cmd.show("sticks","molecule"); cmd.show("spheres","molecule"); cmd.set("sphere_scale",0.25,"molecule")
    cmd.set("stick_radius",0.1,"molecule"); cmd.set("ray_opaque_background",1); cmd.set("ray_shadows","off")
    cmd.set("specular","off"); cmd.set("light_count",3); cmd.set("label_size",22); cmd.set("label_color","black","molecule")
    cmd.label("molecule","elem"); cmd.zoom("molecule",buffer=1); cmd.bg_color("white")
    cmd.png(str(out),width=1200,height=1200,dpi=300,ray=1)
    try: xyz_path.unlink()
    except Exception: pass
    return out


def render_xyz_png(xyz: str, output: str | Path, renderer: str = "pymol") -> Path:
    if renderer == "pymol":
        return render_xyz_png_pymol(xyz, output)
    return render_xyz_png_qt(xyz, output)


def _display_length(value) -> int:
    if value is None:
        return 0
    text = str(value)
    return max((len(part) for part in text.splitlines()), default=0)


def _auto_fit_worksheet_columns(ws, image_headers: set[str] | None = None) -> None:
    """Auto-fit all non-image columns to the widest displayed cell.

    The two structure-image columns remain fixed so inserted PNG images keep a
    consistent physical size.  A modest upper bound prevents an unusually long
    path or warning string from making the worksheet unusably wide.
    """
    image_headers = image_headers or set()
    for column_index in range(1, ws.max_column + 1):
        letter = get_column_letter(column_index)
        header = ws.cell(1, column_index).value
        if column_index == 1:
            ws.column_dimensions[letter].width = 3
            continue
        if header in image_headers:
            ws.column_dimensions[letter].width = 20
            continue
        maximum = 0
        for row_index in range(1, ws.max_row + 1):
            maximum = max(maximum, _display_length(ws.cell(row_index, column_index).value))
        ws.column_dimensions[letter].width = min(max(maximum + 2, 8), 60)


def export_excel_with_structures(
    out_path: str | Path,
    summaries: Sequence[LogSummary],
    include_init: bool = True,
    include_opt: bool = True,
    renderer: str = "pymol",
    progress=None,
) -> Path:
    if not include_init and not include_opt: raise ValueError("Select at least one structure type.")
    wb=Workbook(); ws=wb.active; ws.title="Results"
    headers=["","Cluster"]
    if include_init: headers.append("Init Struc.")
    if include_opt: headers.append("Opted Struc.")
    headers += ["Sym.","State","HF","ZPE","Etot","ΔE","NImag","Freq Done","Opt Conv.","Freq Conv.","Stable1","Stable2"]
    ws.append(headers)
    head_font=Font(name="Times New Roman",size=11,bold=True); body_font=Font(name="Times New Roman",size=10)
    for cell in ws[1]: cell.font=head_font; cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    base_etot=minimum_etot(list(summaries))
    with tempfile.TemporaryDirectory(prefix="chemauto_struct_") as td:
        td=Path(td)
        for idx,s in enumerate(summaries,2):
            row=["",s.file]
            structure_cells=[]
            if include_init: structure_cells.append((len(row)+1,s.init_xyz,"init")); row.append("")
            if include_opt: structure_cells.append((len(row)+1,s.opt_xyz,"opt")); row.append("")
            row += [s.sym,s.state,s.hf,s.zpe,s.etot,delta_e_ev(s,base_etot),s.nimag,s.freq_completed,
                    s.opt_conv.to_compact() if s.opt_conv else "",s.freq_conv.to_compact() if s.freq_conv else "",s.stable_1,s.stable_2]
            ws.append(row); ws.row_dimensions[idx].height=105
            for c in range(1,len(headers)+1): ws.cell(idx,c).font=body_font; ws.cell(idx,c).alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
            for col,xyz,kind in structure_cells:
                if not xyz: ws.cell(idx,col).value="N/A"; continue
                png=td/f"{idx}_{kind}.png"; render_xyz_png(xyz,png,renderer=renderer)
                img=XLImage(str(png)); img.width=105; img.height=105; ws.add_image(img,f"{get_column_letter(col)}{idx}")
            if progress: progress(idx-1,len(summaries),s.file)
        for r in range(2,ws.max_row+1):
            for c in range(1,ws.max_column+1):
                title=ws.cell(1,c).value
                if title in ("HF","ZPE","Etot"): ws.cell(r,c).number_format="0.00000000"
                elif title=="ΔE": ws.cell(r,c).number_format="0.000000"
        _auto_fit_worksheet_columns(ws, image_headers={"Init Struc.", "Opted Struc."})
        wb.save(out_path)
    return Path(out_path)
