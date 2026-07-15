from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from chemauto.core.dos_pairs import identify_dos_pairs, scan_dos_pairs
from chemauto.core.gjf_template import apply_structure_to_template
from chemauto.core.vde_ade import calculate_three_file_vde_ade


class Phase3CoreTests(unittest.TestCase):
    def test_dos_pairing_uses_strict_suffixes(self):
        pairs = identify_dos_pairs([
            "abc_curve.txt", "abc_line.txt", "wrongcurve.txt", "abc_lines.txt",
            "other_curve.txt",
        ])
        self.assertEqual([p.key for p in pairs], ["abc", "other"])
        self.assertTrue(pairs[0].complete)
        self.assertFalse(pairs[1].complete)

    def test_dos_scan_does_not_cross_pair_subfolders(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            (root / "a").mkdir(); (root / "b").mkdir()
            (root / "a" / "same_curve.txt").write_text("", encoding="utf-8")
            (root / "b" / "same_line.txt").write_text("", encoding="utf-8")
            pairs = scan_dos_pairs(root)
            self.assertEqual(len(pairs), 2)
            self.assertTrue(all(not p.complete for p in pairs))

    def test_gjf_template_replaces_first_molecule_and_chk(self):
        template = "%mem=8GB\n%chk=old.chk\n#p upbepbe/def2tzvp opt\n\nTemplate\n\n0 1\nH 0 0 0\n\n--Link1--\n%oldchk=old.chk\n#p geom=allcheck guess=read stable\n\nStable\n\n"
        result = apply_structure_to_template(
            template,
            ["Fe 0.0 0.0 0.0", "C 1.0 0.0 0.0"],
            -1, 10, "newname", "sample.log"
        )
        self.assertIn("%chk=newname.chk", result)
        self.assertIn("-1 10\nFe 0.0 0.0 0.0\nC 1.0 0.0 0.0", result)
        self.assertIn("--Link1--", result)
        self.assertIn("stable", result.lower())

    def test_excel_structure_columns_shift_data_safely(self):
        from openpyxl import load_workbook
        from chemauto.core.gaussian_parser import parse_gaussian_log
        from chemauto.core.excel_structures import export_excel_with_structures
        data = Path(__file__).resolve().parents[1] / "test_data" / "logs"
        record = parse_gaussian_log(str(data / "fe3c2-2A.log"))
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "structures.xlsx"
            export_excel_with_structures(out, [record], include_init=True, include_opt=True, renderer="built-in")
            wb = load_workbook(out)
            ws = wb.active
            headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
            self.assertEqual(headers[:6], [None, "Cluster", "Init Struc.", "Opted Struc.", "Sym.", "State"])
            self.assertEqual(len(ws._images), 2)

    def test_excel_non_image_columns_auto_fit(self):
        from openpyxl import load_workbook
        from chemauto.core.gaussian_parser import parse_gaussian_log
        from chemauto.core.excel_structures import export_excel_with_structures
        data = Path(__file__).resolve().parents[1] / "test_data" / "logs"
        record = parse_gaussian_log(str(data / "fe3c2-2A.log"))
        record.file = "a_very_long_cluster_filename_for_width_check.log"
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "widths.xlsx"
            export_excel_with_structures(out, [record], include_init=True, include_opt=True, renderer="built-in")
            ws = load_workbook(out).active
            self.assertEqual(ws.column_dimensions["C"].width, 20)
            self.assertEqual(ws.column_dimensions["D"].width, 20)
            self.assertGreater(ws.column_dimensions["B"].width, 30)
            self.assertGreaterEqual(ws.column_dimensions["G"].width, len("-3866.59761359") + 2)

    def test_pymol_is_default_excel_renderer(self):
        import inspect
        from chemauto.core.excel_structures import export_excel_with_structures
        signature = inspect.signature(export_excel_with_structures)
        self.assertEqual(signature.parameters["renderer"].default, "pymol")

    def test_vde_ade_three_file_outputs_two_channels(self):
        data = Path(__file__).resolve().parents[1] / "test_data" / "logs"
        anion, results, warnings = calculate_three_file_vde_ade(
            data / "fe3c2-2A.log", data / "fe3c3-3A.log", data / "fe3c4-4A.log"
        )
        self.assertEqual(len(results), 2)
        self.assertEqual([r.label for r in results], ["N−1", "N+1"])
        self.assertIsNotNone(anion.hf)


if __name__ == "__main__":
    unittest.main()
